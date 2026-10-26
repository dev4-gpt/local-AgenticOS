#!/usr/bin/env python3
"""Populate the retained System Design Word template for AgentOS."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


def set_cell(cell, value: str) -> None:
    cell.text = value


def fill_table(table, rows: list[list[str]]) -> None:
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            if row_index < len(table.rows) and col_index < len(table.columns):
                set_cell(table.cell(row_index, col_index), value)


def set_paragraph(doc: Document, prefix: str, value: str) -> None:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            paragraph.text = value
            return
    raise ValueError(f"Template paragraph not found: {prefix}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()

    doc = Document(args.template)
    doc.paragraphs[8].text = "AgentOS | Veloce AgenticOS"
    doc.paragraphs[9].text = "Evidence-Based Production Readiness"

    fill_table(doc.tables[0], [[
        "STATUS\nProposed design",
        "",
        "OWNER\nAryaman",
        "",
        "LAST UPDATED\nOctober 15, 2026",
    ]])
    fill_table(doc.tables[1], [
        ["Authors", "AgentOS project team"],
        ["Reviewers", "Product, runtime/security, operations/billing (owners to be named)"],
        ["Related docs", "Execution board; canonical production topology; production-readiness program"],
        ["Scope", "Architecture and evidence-led sequence to verify AgentOS readiness without unapproved hosted changes or spend."],
    ])

    set_paragraph(
        doc,
        "[Summarize the proposed system",
        "AgentOS is a governed AI workspace with a public intake surface, tenant-scoped workspaces, asynchronous agent jobs, provider controls, artifacts, memory/graph capabilities, and operator tooling. The intended workflow is approved request and sources -> identity, tenant, policy, and budget checks -> bounded job -> source-linked artifact -> human review -> approved use/export.",
    )
    set_paragraph(
        doc,
        "[Describe the workloads",
        "This design organizes the existing stack and its readiness evidence; it does not approve production rollout. Local tests, mock journeys, partial archives, and synthetic personas remain distinct from hosted acceptance and real buyer validation. Hosted mutation, provider spend, public promotion, and destructive data actions require separate approval.",
    )

    fill_table(doc.tables[2], [
        ["Goals", "Non-goals"],
        ["Verify one user journey from sign-in through reviewed, source-linked deliverable and export.", "Claim broad enterprise readiness or market demand from implementation activity."],
        ["Keep jobs tenant-scoped, policy-gated, budget-bounded, traceable, and recoverable.", "Replace the current application, database, cloud, or worker stack without evidence."],
        ["Define source-pack coverage, provenance, and graph-import acceptance.", "Use session cookies or bypass platform controls for Skool collection."],
        ["Bind every readiness claim to a source candidate, named environment, and timestamped evidence.", "Treat local tests, synthetic studies, partial imports, or historical deployments as production proof."],
    ])

    set_paragraph(
        doc,
        "[Describe the current state",
        "The repository has a Next.js client runtime, public Vercel intake/editorial surface, Supabase-backed tenant state, Cloud Tasks and private workers, provider adapters, private artifact storage, graph/notebook workflows, observability, and a release verifier. The topology record distinguishes the public site from the authenticated workspace and warns that shared Supabase project migrations affect more than the beta tenant.",
    )
    set_paragraph(
        doc,
        "[Explain the proposed system boundary",
        "The readiness gap is evidence reconciliation, not simply missing tools. Fixture-backed UI tests are not a signed-in customer journey; local configuration is not hosted configuration; a partial archive is not a complete import; synthetic personas are not buyers; and a deployed image without source identity is not a reproducible release. Invariant: one immutable candidate, named environment/tenant, bounded action, timestamped result, and evidence artifact per readiness claim.",
    )
    set_paragraph(
        doc,
        "Figure 1.",
        "Figure 1. Logical flow: Member browser -> Public Vercel information/intake or authenticated tenant runtime -> Supabase tenant state -> Cloud Tasks -> private worker -> policy/budget-gated provider -> private artifact proxy -> human review. Operator/local tools remain a separate trust boundary. This is a target boundary, not proof of current deployment parity.",
    )

    fill_table(doc.tables[3], [
        ["Component", "Responsibility", "Primary storage", "Failure behavior"],
        ["Public surface", "Product information and intake; never impersonates workspace login.", "Vercel and configured intake state", "Fail closed on unavailable/misrouted login; no legacy fallback."],
        ["Tenant runtime", "Authenticate, resolve membership, enforce capability and tenant-scoped APIs.", "Supabase Auth and tenant data", "Deny unauthenticated, inactive, or cross-tenant requests."],
        ["Policy/budget boundary", "Check capability, consent, limits, approval, and reservation before side effects.", "Versioned config and durable ledger", "Fail closed when critical policy or budget state is unavailable."],
        ["Dispatcher/worker", "Claim bounded jobs, call adapters, persist outputs, settle/refund.", "Cloud Tasks, job/ledger state, secrets", "Idempotent retry only; reconcile uncertain external outcomes."],
        ["Source/artifact pipeline", "Validate approved sources, preserve provenance, store private artifacts, mediate review/export.", "Private object store, tenant DB/graph", "Reject invalid/unauthorized packs; never publish partial as complete."],
    ])

    flow = [
        "1. Resolve an authenticated member and active tenant from server-side session state.",
        "2. Validate capability, input schema, consent, source ownership, idempotency, and limits; do not trust client document text over approved source IDs.",
        "3. Load effective policy/config and required approval/budget state. Unknown critical state blocks execution.",
        "4. Persist the job and budget reservation before any provider or external-tool side effect.",
        "5. A private worker claims the job, rechecks eligibility, uses minimum required secrets, and records bounded/redacted stage outcomes.",
        "6. Persist private artifacts with source provenance; settle/refund budget; transition to a terminal state without duplicate effects on retry.",
        "7. Show durable status and sources; require human review before policy-controlled export or consequential action.",
    ]
    for prefix, value in zip(
        ["[Describe how a request", "[Describe validation", "[Describe which configuration", "[Describe the primary decision", "[Describe which durable state", "[Describe downstream calls", "[Describe final state"],
        flow,
    ):
        set_paragraph(doc, prefix, value)

    fill_table(doc.tables[4], [
        ["Field", "Type", "Required", "Description"],
        ["tenant identity", "Opaque ID", "Yes", "Resolved from trusted runtime context; never body-only authority."],
        ["actor identity", "Authenticated subject", "Yes", "Member or service responsible for action."],
        ["operation identity", "Stable ID", "Async work", "Supports idempotency, retry, audit correlation."],
        ["source references", "IDs + versions + URLs", "When grounded", "Must resolve within same tenant and be approved for this operation."],
        ["policy/config version", "Version or hash", "Gated work", "Explains the allowlist, budget, capability, and approval applied."],
        ["lifecycle status", "Versioned enum", "Yes", "Validated transitions; terminal states do not silently reopen."],
        ["budget/provenance", "Ledger refs + source/review metadata", "When applicable", "Reconcile spend and preserve source and human-review context."],
    ])
    set_paragraph(doc, "[State the ordering", "Authorization and policy checks precede enqueue and storage access; durable job and reservation precede external side effects.")
    set_paragraph(doc, "[State how writes", "Use stable operation and idempotency identities across request, worker claim, provider adapter, artifact write, and settlement.")
    set_paragraph(doc, "[State which versions", "Capture candidate/source/config/policy/schema references and event time needed to audit the decision and safely replay it.")
    set_paragraph(doc, "[State what this data", "The durable database ledger/job state is authoritative for permissions, budget, and status; telemetry is diagnostic, not a financial or authorization source of truth.")
    set_paragraph(doc, "[Link to interface", "Update the existing versioned route/schema contracts and compatibility tests with each contract release; do not invent a parallel schema.")

    set_paragraph(doc, "[Explain the consistency", "Queue delivery is treated as at-least-once. Use durable claim/lease and idempotency patterns; do not claim exactly-once provider execution unless verified. Replays are operator-controlled, tenant-scoped, and must not re-charge, overwrite approved output, cross tenants, or bypass consent.")
    fill_table(doc.tables[5], [
        ["Scenario", "Expected behavior", "Reasoning"],
        ["Duplicate request/delivery", "Continue existing operation; no duplicate reservation/artifact.", "Avoid duplicate work and spend."],
        ["Reservation/job write fails", "Do not call provider; surface safe retryable failure.", "No external side effect without durable admission."],
        ["Provider timeout with uncertain completion", "Reconcile before retry unless provider idempotency is verified.", "Avoid duplicate charges and inconsistent output."],
        ["Policy changes in flight", "Use recorded version, with explicit cancel/recheck for revocation.", "Keep decisions auditable and deterministic."],
    ])

    for prefix, value in zip(
        ["[Describe authentication", "[Describe data minimization", "[Describe credential", "[Describe safe defaults", "[Describe retention"],
        [
            "Authenticate every client route; resolve active membership server-side; enforce RLS and API authorization independently. Add two-tenant negative tests for every new access path.",
            "Treat documents, URLs, graph imports, retrieved text, and model output as untrusted. Bound parser/input sizes and prevent secrets, raw customer text, or full prompts/results from entering logs and fixtures.",
            "Keep provider/signing keys server-side in the existing secret boundary. Artifacts remain private and are served through an authenticated tenant-checked proxy; do not broaden IAM for convenience.",
            "Migrations, replay, admin/debug access, secret rotation, retention, IAM, traffic shifts, and paid provider use require explicit scope, preflight, evidence, and rollback. Billing alerts are not hard spend caps.",
            "The owner selected 28 active days from item creation plus a separate seven-day object recovery window (up to 35 days recoverable). Inventory every tenant data class and reconcile database, exports, backups, and vendor copies; financial/credit and minimized security/audit evidence require separately approved schedules.",
        ],
    ):
        set_paragraph(doc, prefix, value)

    fill_table(doc.tables[6], [
        ["Signal", "SLO or alert", "Owner", "Launch gate"],
        ["Customer journey", "Track step completion/failure; baseline and numeric SLO to be approved.", "Product/runtime owner TBD", "Required"],
        ["Tenant/policy denials", "Test expected denials; alert on unexpected cross-tenant access or bypass.", "Security/runtime owner TBD", "Required"],
        ["Job/provider lifecycle", "Reconcile queue, timeout, retry, settlement/refund, terminal state.", "Operations owner TBD", "Required"],
        ["Cost and retention", "Verify hard provider caps, ledger reconciliation, per-item expiry and recovery semantics.", "Billing/privacy owners TBD", "Required"],
        ["Build identity and telemetry", "Record source/image/config/migration/rollback; test bounded redacted trace and alert delivery.", "Release/observability owners TBD", "Required"],
        ["Rollout constraint", "Controlled beta first; read-only checks -> local tests -> separately approved hosted test -> verified rollback -> release review.", "No public promotion implied", "Required"],
    ])

    fill_table(doc.tables[7], [
        ["Alternative", "Why it was considered", "Why it was not selected"],
        ["Replace current cloud/app stack", "Might simplify greenfield architecture.", "Adds migration and operational risk before proving existing stack is blocker."],
        ["Add a new observability platform now", "Dedicated trace UI may help later.", "Existing local contracts need hosted evidence first; new infra adds cost and privacy work."],
        ["Create another tracker", "Could appear to add structure.", "Would split the canonical execution board and evidence source of truth."],
        ["Scale to 10,000 synthetic buyers", "Broader scenario generation.", "Current small study failed quality gate; scale cannot establish real willingness to pay."],
    ])

    for prefix, value in zip(
        ["[Open question 1", "[Open question 2", "[Open question 3", "[Open question 4"],
        [
            "Which signed-in Veloce Chrome profile/tab can be connected? Current available browser inventory exposes only the unrelated College profile.",
            "Which exact service, revision, tenant, and Supabase project are authorized for the next read-only parity check? Beta shares a database project with production.",
            "Who is the billing reviewer for the separate financial/credit schedule, and what backup/vendor deletion terms apply? The selected content policy is 28 active days plus seven days of object recovery; local code does not prove complete coverage or hosted enforcement.",
            "Who owns operations, privacy/legal, billing, rollback, support, and the missing archive exports? Buyer segment/deliverable frequency also needs real interviews.",
        ],
    ):
        set_paragraph(doc, prefix, value)

    set_paragraph(
        doc,
        "[State the recommended decision",
        "Adopt this document as the architecture companion to the existing execution board; do not add a replacement infrastructure stack. Keep one active critical-path implementation item, parallelizing only independent read-only work. Begin with evidence-baseline reconciliation, then local journey, source archive/import readiness, hosted parity, offline AI/trace gates, and finally the strict release decision. No hosted mutation, paid request, or public launch is authorized by this design.",
    )
    fill_table(doc.tables[8], [
        ["Milestone", "Deliverable", "Exit criteria"],
        ["M0: Evidence baseline", "Candidate SHA, clean state, verifier, test receipts, environment identities.", "One current board summary; stale evidence labeled; no guessed percentage."],
        ["M1: Local journey", "Supported-runtime sign-in-to-reviewed-deliverable journey and negative tests.", "Journey/regressions pass; mocks clearly distinguished from real-user evidence."],
        ["M2: Archive/import", "Canonical coverage, validated user-supplied pack, dry-run and scoped import plan.", "All in-scope content accounted; discrepancies closed; idempotent tenant-scoped import."],
        ["M3: Hosted parity", "Same-candidate, same-user/tenant read-only comparison; config and identity evidence.", "Differences resolved or explicit; isolation/storage/retention checked before migrations."],
        ["M4-M5: AI and release", "Offline eval/trace gates, required approvals, complete evidence bundle and rollback.", "Strict verifier ready for exact candidate; otherwise record not-ready and missing gates."],
    ])

    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)


if __name__ == "__main__":
    main()
